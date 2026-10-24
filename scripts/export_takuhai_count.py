# -*- coding: utf-8 -*-
"""店舗で受け付けた宅配の件数（＝送料ベースの発送件数）を店舗別・月次でExcel出力。

件数の定義: 商品カテゴリ「送料」の数量合計（発送1件につき送料1、を件数とみなす）。
事業年度 = 9月〜翌8月。FY2023(2022/9〜)〜進行中FY2026。
"""
import os
from datetime import date
from collections import defaultdict

from dotenv import load_dotenv
from supabase import create_client
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
load_dotenv(os.path.join(BASE, ".env"))
sb = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_ROLE_KEY"])

OUT_DIR = "/Users/tanakatsuyoshi/Desktop/餃子/04_経営管理/★取締役会/第18回202607/販売記録"
OUT_PATH = os.path.join(OUT_DIR, "宅配発送件数_店舗別月次_FY2023-2026.xlsx")
START, END = "2022-09-01", "2026-09-01"


def fiscal_year(d: date) -> int:
    return d.year + 1 if d.month >= 9 else d.year


def parse(s):
    return date.fromisoformat(s)


def month_str(d):
    return f"{d.year}-{d.month:02d}"


def months_between(mmin, mmax):
    y, m = int(mmin[:4]), int(mmin[5:7])
    ye, me = int(mmax[:4]), int(mmax[5:7])
    out = []
    while (y, m) <= (ye, me):
        out.append(f"{y}-{m:02d}")
        m += 1
        if m > 12:
            m, y = 1, y + 1
    return out


def fetch_all(table, columns, filters):
    """order=id 固定でページング（行の重複・欠落を防止）。"""
    rows, step, start = [], 1000, 0
    while True:
        q = sb.table(table).select(columns).order("id").range(start, start + step - 1)
        for f in filters:
            q = f(q)
        batch = q.execute().data or []
        rows.extend(batch)
        if len(batch) < step:
            break
        start += step
    return rows


print("店舗マスタ取得...")
segs = {s["id"]: (s["code"], s["name"])
        for s in sb.table("segments").select("id,code,name").execute().data}

print("送料(宅配)データ取得...")
rows = fetch_all(
    "product_sales", "segment_id,sale_date,product_category_name,quantity",
    [lambda q: q.eq("product_category_name", "送料"),
     lambda q: q.gte("sale_date", START), lambda q: q.lt("sale_date", END)],
)

# 集計: (code,name) -> month -> 件数
grid = defaultdict(lambda: defaultdict(float))
present = set()
for r in rows:
    code, name = segs.get(r["segment_id"], ("", r["segment_id"]))
    mth = month_str(parse(r["sale_date"]))
    grid[(code, name)][mth] += float(r.get("quantity") or 0)
    present.add(mth)

months = months_between(min(present), max(present))

# 列プラン: 月 + 各年度末にFY計 + 全期間計
plan, prev = [], None
for mth in months:
    fy = fiscal_year(parse(mth + "-01"))
    if prev is not None and fy != prev:
        plan.append(("fyt", prev))
    plan.append(("m", mth))
    prev = fy
if prev is not None:
    plan.append(("fyt", prev))
plan.append(("gt", None))


def plan_values(cm):
    out = []
    for kind, key in plan:
        if kind == "m":
            out.append(cm.get(key))
        elif kind == "fyt":
            s = sum(v for mm, v in cm.items() if fiscal_year(parse(mm + "-01")) == key)
            out.append(s or None)
        else:
            out.append(sum(cm.values()) or None)
    return out


# ---- Excel ----
wb = Workbook()
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
TOT_FILL = PatternFill("solid", fgColor="FCE4D6")
GRAND_FILL = PatternFill("solid", fgColor="C6E0B4")
LABEL_FILL = PatternFill("solid", fgColor="DDEBF7")
LABEL_FONT = Font(bold=True, size=10)
BOLD = Font(bold=True, size=10)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# Sheet1: 説明
ws = wb.active
ws.title = "説明"
ws["A1"] = "宅配発送件数（店舗別・月次）"
ws["A1"].font = Font(bold=True, size=14)
for i, t in enumerate([
    "出力元: KPIダッシュボード（MaruokaKPI / Supabase）",
    f"出力日: {date.today().isoformat()}",
    "事業年度: 9月〜翌8月（例: FY2024 = 2023年9月〜2024年8月）",
    "",
    "【件数の定義】",
    "  ・店舗で受け付けた宅配の件数 ＝ 商品カテゴリ「送料」の数量合計（発送1件につき送料1）。",
    "  ・POS商品別実績（product_sales）の送料行を店舗×月で集計。",
    "  ・値引・返品等により月によっては送料数量が0となる場合あり。",
    "",
    "【シート】",
    "  ■ 宅配件数_店舗別月次 : 縦=店舗、横=月＋年度計(FYxxxx 計)＋全期間計。末尾に全店合計。",
    "  ■ 明細(縦持ち)        : 店舗×月の件数（ピボット用）。",
    "",
    "【対象事業年度】",
    "  FY2023(2022/9〜2023/8) … 送料POSは2022/12以降のみ収録",
    "  FY2024・FY2025 … 通期 / FY2026 … 進行中（最新実績月まで）",
], 3):
    ws[f"A{i}"] = t
autosize(ws, [90])

# Sheet2: マトリクス
ws = wb.create_sheet("宅配件数_店舗別月次")
nlab = 2
fy_row, mo_row = [""] * nlab, ["店舗コード", "店舗名"]
for kind, key in plan:
    if kind == "m":
        fy_row.append(f"FY{fiscal_year(parse(key + '-01'))}")
        mo_row.append(key)
    elif kind == "fyt":
        fy_row.append(f"FY{key}")
        mo_row.append(f"FY{key} 計")
    else:
        fy_row.append("全期間")
        mo_row.append("全期間 計")
ws.append(fy_row)
ws.append(mo_row)
ncol = nlab + len(plan)
style_header(ws, 1, ncol)
style_header(ws, 2, ncol)


def store_key(k):
    code = k[0]
    return (0, int(code)) if str(code).isdigit() else (1, str(code))


grand = defaultdict(float)
for k in sorted(grid, key=store_key):
    code, name = k
    ws.append([code, name] + plan_values(grid[k]))
    r = ws.max_row
    for c in range(1, nlab + 1):
        ws.cell(r, c).fill = LABEL_FILL
        ws.cell(r, c).font = LABEL_FONT
    for i, (kind, _) in enumerate(plan):
        cell = ws.cell(r, nlab + 1 + i)
        cell.number_format = '#,##0'
        if kind in ("fyt", "gt"):
            cell.fill = TOT_FILL
            cell.font = BOLD
    for mth, v in grid[k].items():
        grand[mth] += v
# 全店合計
ws.append(["", "【全店】合計"] + plan_values(grand))
r = ws.max_row
for c in range(1, ncol + 1):
    ws.cell(r, c).fill = GRAND_FILL
    ws.cell(r, c).font = BOLD
    if c > nlab:
        ws.cell(r, c).number_format = '#,##0'
ws.freeze_panes = ws.cell(row=3, column=nlab + 1).coordinate
autosize(ws, [10, 14] + [12 if k in ("fyt", "gt") else 9 for k, _ in plan])

# Sheet3: 明細（縦持ち）
ws = wb.create_sheet("明細(縦持ち)")
ws.append(["事業年度", "年月", "店舗コード", "店舗名", "宅配件数(発送)"])
style_header(ws, 1, 5)
ws.freeze_panes = "A2"
detail = []
for k in sorted(grid, key=store_key):
    code, name = k
    for mth in months:
        v = grid[k].get(mth)
        if v:
            d = parse(mth + "-01")
            detail.append((f"FY{fiscal_year(d)}", mth, code, name, v))
detail.sort(key=lambda x: (x[0], x[1], store_key((x[2],))))
for row in detail:
    ws.append(row)
    ws.cell(ws.max_row, 5).number_format = '#,##0'
autosize(ws, [10, 9, 10, 14, 14])

os.makedirs(OUT_DIR, exist_ok=True)
wb.save(OUT_PATH)
print("保存:", OUT_PATH)
print(f"  店舗数:{len(grid)}  月数:{len(months)}  明細:{len(detail)}行")
print(f"  全店 宅配件数 全期間計 = {sum(grand.values()):,.0f}")
